import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field

from openpyxl import load_workbook
from pypdf import PdfReader

FIELDS = (
    "tracking_code",
    "address",
    "city",
    "state",
    "postal_code",
    "recipient_reference",
    "delivery_notes",
)

ALIASES = {
    "tracking_code": {"tracking_code", "codigo", "cod", "rastreio", "tracking", "pedido", "pacote"},
    "address": {"address", "endereco", "logradouro", "rua", "endereco_completo"},
    "city": {"city", "cidade", "municipio"},
    "state": {"state", "uf", "estado"},
    "postal_code": {"postal_code", "cep", "codigo_postal"},
    "recipient_reference": {"recipient_reference", "destinatario", "cliente", "referencia"},
    "delivery_notes": {"delivery_notes", "observacao", "observacoes", "obs", "notas"},
}

XLSX_MAGIC = b"PK\x03\x04"
PDF_MAGIC = b"%PDF"

# Padrões de romaneio/lista de entrega em PDF (texto livre, sem colunas garantidas).
_STREET = re.compile(
    r"^(?:(?:rua|avenida|av|estrada|rodovia|rod|travessa|trav|tv|alameda|al|pra[cç]a|"
    r"p[cç]a|largo|lgo|beco|quadra|qd|viela|caminho|boulevard|blvd|setor|loteamento|"
    r"condom[ií]nio|ladeira|marginal|servid[aã]o)\b|(?:r|av|estr|rod|trav|tv|al|p[cç]a|"
    r"lgo|bc|qd|vl|blvd)\.)\s*[\wÀ-ÿ]",
    re.IGNORECASE,
)
# UF só em maiúsculas, para não confundir "Pa"/"Al" do fim de um nome com estado.
_UFS = "(?-i:AC|AL|AP|AM|BA|CE|DF|ES|GO|MA|MT|MS|MG|PA|PB|PR|PE|PI|RJ|RN|RS|RO|RR|SC|SP|SE|TO)"
_CITY = r"[A-Za-zÀ-ÖØ-öø-ÿ'][A-Za-zÀ-ÖØ-öø-ÿ' .]*?"
_CEP = re.compile(r"\bcep\b[:\s.\-]*(\d{5}-?\d{3})|\b(\d{5}-\d{3})\b", re.IGNORECASE)
_CEP_ANY = re.compile(r"(?:\bcep\b[:\s.\-]*)?\b(\d{5}-?\d{3})\b", re.IGNORECASE)
# "..., Osasco/SP" | "... - Osasco - SP" | "..., Osasco, SP"
_CITY_UF_TAIL = re.compile(rf"[-,]\s*({_CITY})\s*(?:/|\s-\s|,)\s*({_UFS})\s*$", re.IGNORECASE)
# linha de tabela sem separadores: "..., 136 Osasco SP"
_CITY_UF_BARE_TAIL = re.compile(rf"(\d+[A-Za-z]?|s/?n)\s+({_CITY})\s+({_UFS})\s*$", re.IGNORECASE)
_CITY_UF_LINE = re.compile(rf"^({_CITY})\s*[-/,]\s*({_UFS})$", re.IGNORECASE)
# tabela "Rua X 136 Bairro - Cidade": número e bairro/cidade após o logradouro,
# sem separador e sem UF. O último número da linha é o da casa/prédio.
_NUM_PLACE_TAIL = re.compile(r"\b(\d{1,6}[A-Za-z]?|s/?n)\s+([^\d]{2,})$", re.IGNORECASE)
# célula de tabela: só número ("136", "s/n") ou só nome de lugar ("Osasco")
_BARE_NUMBER = re.compile(r"^(?:s/?n|\d{1,6}[A-Za-z]?)$", re.IGNORECASE)
# número de linha da tabela ("1", "23"), usado como separador de registro
_ROW_INDEX = re.compile(r"^\s*(?:n[ºo°]\s*)?\d{1,4}\s*$", re.IGNORECASE)
_BARE_PLACE = re.compile(rf"^({_CITY})(?:\s*[-–—]\s*({_CITY}))?$", re.IGNORECASE)
# palavras de cabeçalho de tabela ("Nº", "Rua / Avenida", "Número", "CEP"...)
_HEADER_WORDS = {
    "n",
    "no",
    "numero",
    "item",
    "ordem",
    "seq",
    "sequencia",
    "rua",
    "avenida",
    "av",
    "logradouro",
    "endereco",
    "bairro",
    "cidade",
    "municipio",
    "cep",
    "uf",
    "estado",
    "nome",
    "cliente",
    "destinatario",
    "recebedor",
    "codigo",
    "rastreio",
    "tracking",
    "obs",
    "observacao",
    "observacoes",
    "nota",
    "notas",
    "status",
    "data",
    "pacote",
    "quantidade",
    "qtd",
    "assinatura",
    "complemento",
    "referencia",
    "telefone",
    "doc",
    "documento",
    "num",
}


def _is_header_line(text: str) -> bool:
    tokens = [t for t in normalize_header(text).split("_") if t]
    return bool(tokens) and all(t in _HEADER_WORDS for t in tokens)


def _attach_place(rec: dict[str, str | None], place: str) -> None:
    """'Planalto Paulista - São Paulo' -> cidade 'São Paulo' e bairro no endereço."""
    parts = [p.strip(" .,") for p in re.split(r"\s*[-–—/]\s*", place) if p.strip(" .,")]
    if not parts:
        return
    rec["city"] = parts[-1]
    if len(parts) > 1:
        rec["address"] = f"{rec['address']} - {' - '.join(parts[:-1])}"


_LIST_NUMBER = re.compile(r"^\s*(?:n[ºo°]\s*)?\d{1,4}\s*(?:[.)\-:]\s*|\s+)(?=\D)")
_ADDRESS_LABEL = re.compile(
    r"^(?:endere[cç]o(?:\s+de\s+entrega)?|end\.|destino|logradouro|entregar\s+em)\s*[:\-]\s*",
    re.IGNORECASE,
)
_LEADING_CODE = re.compile(r"^([A-Z]{2,}[\w\-]*\d{6,}[A-Z]*|\d{9,})\s+(?=\D)", re.IGNORECASE)
_LABELED_CODE = re.compile(
    r"(?:pedido|pacote|c[oó]d(?:igo)?\.?|rastreio|tracking|ordem)\b"
    r"\s*[:#nºo°.\-]*\s*([A-Z0-9][\w\-]{5,})",
    re.IGNORECASE,
)
_BARE_CODE = re.compile(
    r"^(?!\d{5}-?\d{3}$)(?=.{8,}$)[A-Z0-9][A-Z0-9\-]*\d[A-Z0-9\-]*$", re.IGNORECASE
)
_DASHES = re.compile(r"\s*[–—]\s*")
_NOTE = re.compile(r"^obs(?:erva[çc][ãa]o)?\.?\s*[:.\-]?\s*(.+)", re.IGNORECASE)


class ImportFileError(ValueError):
    pass


@dataclass
class ParsedImport:
    rows: list[dict[str, str | None]] = field(default_factory=list)
    errors: list[dict] = field(default_factory=list)
    mapping: dict[str, str] = field(default_factory=dict)


def normalize_header(value: str) -> str:
    text = unicodedata.normalize("NFKD", str(value)).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", "_", text.strip().lower()).strip("_")


def normalize_postal_code(value: str) -> str | None:
    digits = re.sub(r"\D", "", value)
    return f"{digits[:5]}-{digits[5:]}" if len(digits) == 8 else None


def _read_table(filename: str, content: bytes) -> list[list[str]]:
    is_zip = content.startswith(XLSX_MAGIC)
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "xlsx" and is_zip:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
        return [
            ["" if cell is None else str(cell).strip() for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
    if ext in ("csv", "txt") and not is_zip and b"\x00" not in content[:4096]:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return [[c.strip() for c in row] for row in csv.reader(io.StringIO(text), dialect)]
    raise ImportFileError("Formato não suportado: envie um arquivo CSV, XLSX ou PDF válido")


def split_address_line(line: str) -> dict[str, str | None] | None:
    """Interpreta uma linha de endereço em vários formatos: numeração de lista,
    rótulo "Endereço:", código de rastreio na frente, CEP e cidade/UF na mesma
    linha (com "/", "-", "," ou só espaços), separador ";" e caixa alta.
    Retorna `None` se a linha não começar por um logradouro."""
    text = re.sub(r"\s+", " ", _DASHES.sub(" - ", line.replace(";", ","))).strip()
    text = _ADDRESS_LABEL.sub("", _LIST_NUMBER.sub("", text, count=1), count=1)
    rec: dict[str, str | None] = {}
    code = _LEADING_CODE.match(text)
    if code and _STREET.match(text[code.end() :]):
        rec["tracking_code"], text = code.group(1), text[code.end() :]
    if not _STREET.match(text):
        return None

    cep = _CEP_ANY.search(text)
    if cep:
        rec["postal_code"] = cep.group(1)
        text = (text[: cep.start()] + text[cep.end() :]).strip(" ,-")
    if tail := _CITY_UF_TAIL.search(text):
        rec["city"], rec["state"] = tail.group(1).strip(" ."), tail.group(2).upper()
        text = text[: tail.start()]
    elif bare := _CITY_UF_BARE_TAIL.search(text):
        rec["city"], rec["state"] = bare.group(2).strip(" ."), bare.group(3).upper()
        text = text[: bare.end(1)]
    elif num_place := _NUM_PLACE_TAIL.search(text):
        # linha de tabela: "Rua X 136 Bairro - Cidade" sem separadores nem UF
        rec["address"] = f"{text[: num_place.start()].strip(' ,-')}, {num_place.group(1)}"
        _attach_place(rec, num_place.group(2))
    text = text.strip(" ,-")
    rec["address"] = rec.get("address") or (text.title() if text.isupper() else text)
    if rec.get("city") and rec["city"].isupper():
        rec["city"] = rec["city"].title()
    return rec


def parse_pdf_lines(lines: list[str]) -> list[dict[str, str | None]]:
    """Extrai registros de um romaneio em texto: cada linha de logradouro inicia um
    registro e cada linha vizinha (até 4 de distância) é atribuída ao logradouro
    mais próximo, fornecendo CEP, cidade/UF, código de rastreio e observação."""
    lines = [_DASHES.sub(" - ", line) for line in lines]
    records: dict[int, dict[str, str | None]] = {}
    for i, line in enumerate(lines):
        if (rec := split_address_line(line)) is not None:
            records[i] = rec
    streets = list(records)

    for j, near in enumerate(lines):
        if j in records or _is_header_line(near):
            continue
        # dono = logradouro mais próximo (empate fica com o de cima); o de
        # baixo é descartado se um número de linha da tabela separa os dois
        above = max((s for s in streets if s < j), default=None)
        below = min((s for s in streets if s > j), default=None)
        if below is not None and any(_ROW_INDEX.match(lines[k]) for k in range(j + 1, below)):
            below = None
        owner = min(
            (s for s in (above, below) if s is not None),
            key=lambda s: (abs(s - j), s > j),
            default=None,
        )
        if owner is None or abs(owner - j) > 4:
            continue
        rec = records[owner]
        cep = _CEP.search(near)
        if cep:
            rec.setdefault("postal_code", cep.group(1) or cep.group(2))
            continue
        labeled = _LABELED_CODE.search(near)
        if labeled:
            rec.setdefault("tracking_code", labeled.group(1))
            continue
        city = _CITY_UF_LINE.match(near.strip())
        if city and rec.get("city") is None:
            rec["city"], rec["state"] = city.group(1).strip(" ."), city.group(2).upper()
            continue
        note = _NOTE.match(near)
        if note:
            rec.setdefault("delivery_notes", note.group(1).strip())
            continue
        # tabela célula a célula: "136" e "Osasco" vêm em linhas depois do logradouro
        if (
            owner < j
            and _BARE_NUMBER.match(near)
            and not re.search(r"\d", rec.get("address") or "")
        ):
            rec["address"] = f"{rec['address']}, {near.upper()}"
            continue
        if owner < j and rec.get("city") is None:
            if _BARE_PLACE.match(near):
                _attach_place(rec, near.strip())
                continue
        if _BARE_CODE.match(near):
            rec.setdefault("tracking_code", near)
            continue
    return list(records.values())


def _parse_pdf(content: bytes, max_rows: int) -> ParsedImport:
    try:
        reader = PdfReader(io.BytesIO(content))
        lines = [
            line.strip()
            for page in reader.pages
            for line in (page.extract_text() or "").splitlines()
            if line.strip()
        ]
    except ImportFileError:
        raise
    except Exception as exc:
        raise ImportFileError("PDF inválido ou protegido por senha") from exc
    if not lines:
        raise ImportFileError(
            "PDF sem texto legível — pode ser um documento digitalizado como imagem"
        )
    records = parse_pdf_lines(lines)
    if len(records) > max_rows:
        raise ImportFileError(f"Limite de {max_rows} linhas excedido ({len(records)} endereços)")
    result = ParsedImport(mapping={"source": "pdf"})
    for rec in records:
        if rec.get("postal_code"):
            rec["postal_code"] = normalize_postal_code(rec["postal_code"])
        if rec.get("state"):
            rec["state"] = rec["state"].upper()[:2]
        result.rows.append({f: rec.get(f) for f in FIELDS})
    if not result.rows:
        raise ImportFileError("Nenhum endereço reconhecido no PDF")
    return result


def detect_mapping(headers: list[str]) -> dict[str, str]:
    """Mapeia campo do Entreggo -> cabeçalho original da planilha."""
    mapping = {}
    for header in headers:
        norm = normalize_header(header)
        for field_name, aliases in ALIASES.items():
            if norm in aliases and field_name not in mapping:
                mapping[field_name] = header
    return mapping


def parse_import(
    filename: str, content: bytes, mapping: dict[str, str] | None = None, max_rows: int = 2000
) -> ParsedImport:
    if content.startswith(PDF_MAGIC):
        if mapping:
            raise ImportFileError("Mapeamento de colunas não se aplica a PDF")
        return _parse_pdf(content, max_rows)
    table = [(n, row) for n, row in enumerate(_read_table(filename, content), 1) if any(row)]
    if not table:
        raise ImportFileError("Arquivo vazio")
    headers, data = table[0][1], table[1:]
    if len(data) > max_rows:
        raise ImportFileError(f"Limite de {max_rows} linhas excedido ({len(data)} linhas)")

    mapping = mapping or detect_mapping(headers)
    unknown = {k for k in mapping if k not in FIELDS} | {
        v for v in mapping.values() if v not in headers
    }
    if unknown:
        raise ImportFileError(f"Mapeamento inválido: {sorted(unknown)}")
    if "tracking_code" not in mapping and "address" not in mapping:
        raise ImportFileError("Nenhuma coluna de código ou endereço reconhecida")

    index = {f: headers.index(h) for f, h in mapping.items()}
    result = ParsedImport(mapping=mapping)
    for line_no, raw in data:
        row = {f: (raw[i].strip() or None) if i < len(raw) else None for f, i in index.items()}
        if not row.get("tracking_code") and not row.get("address"):
            result.errors.append({"line": line_no, "message": "Linha sem código e sem endereço"})
            continue
        # Endereço completo numa célula só: separa cidade/UF/CEP que não vieram em colunas.
        if row.get("address") and (split := split_address_line(row["address"])):
            row["address"] = split["address"]
            for key in ("city", "state", "postal_code", "tracking_code"):
                if split.get(key) and not row.get(key):
                    row[key] = split[key]
        if row.get("postal_code"):
            cep = normalize_postal_code(row["postal_code"])
            if cep is None:
                result.errors.append({"line": line_no, "message": "CEP inválido; campo ignorado"})
            row["postal_code"] = cep
        if row.get("state"):
            row["state"] = row["state"].upper()[:2]
        result.rows.append(row)
    return result
